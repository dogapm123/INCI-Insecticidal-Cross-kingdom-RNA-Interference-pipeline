#!/usr/bin/env python3
"""End-to-end dsRNA to siRNA/transcript mismatch scoring workflow.

For each dsRNA FASTA record, this pipeline:

1. Generates all dsRIP-style siRNA duplexes with ``dsRNA_garden_of_edan``.
2. Scores complete siRNA guides with ``dsrip_sirna_feature_api``.
3. Runs the ``MFE_ratio`` RNAplex pipeline against transcript FASTA records.
4. Scores each retained siRNA-transcript interaction with
   ``mismatch_tolerance_score``.
5. Writes per-phase TSV/JSON outputs plus a final Excel workbook.
"""

from __future__ import annotations

import argparse
import csv
import json
import shutil
from dataclasses import asdict
from pathlib import Path
from typing import Any, Iterable, Sequence

from dsRNA_garden_of_edan import (
    DEFAULT_OVERHANG_LENGTH,
    DEFAULT_SIRNA_LENGTH,
    DsRNAInputError,
    DsRNARecord,
    generate_sirna_duplexes,
    read_fasta as read_dsrna_fasta,
)
from dsrip_sirna_feature_api import predict_sirna_features
from MFE_ratio import run_mfe_analysis
from mismatch_tolerance_score import pair_annotations_from_mfe_row, score_file, write_detailed_report


def safe_name(name: str) -> str:
    import re

    return re.sub(r"[^A-Za-z0-9_.-]+", "_", name).strip("_") or "dsRNA"


def read_fasta_records(path: Path, limit: int = 0) -> list[tuple[str, str]]:
    records: list[tuple[str, str]] = []
    current_name: str | None = None
    sequence_parts: list[str] = []

    def finish() -> None:
        nonlocal current_name, sequence_parts
        if current_name is None:
            return
        records.append((current_name, "".join(sequence_parts)))

    with path.open("r", encoding="utf-8") as handle:
        for raw_line in handle:
            line = raw_line.strip()
            if not line:
                continue
            if line.startswith(">"):
                finish()
                if limit and len(records) >= limit:
                    return records
                current_name = line[1:].split()[0]
                if not current_name:
                    raise ValueError(f"Empty FASTA header in {path}")
                sequence_parts = []
            elif current_name is None:
                raise ValueError(f"Sequence found before first FASTA header in {path}")
            else:
                sequence_parts.append(line)
    finish()
    if limit:
        records = records[:limit]
    return records


def write_fasta_records(records: Iterable[tuple[str, str]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for name, sequence in records:
            handle.write(f">{name}\n")
            for start in range(0, len(sequence), 80):
                handle.write(sequence[start : start + 80] + "\n")


def write_tsv(rows: Sequence[dict[str, Any]], path: Path, fieldnames: Sequence[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fieldnames is None:
        fieldnames = collect_fieldnames(rows)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(fieldnames), delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)


def read_tsv(path: Path) -> list[dict[str, Any]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def write_json(data: Any, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2), encoding="utf-8")


def collect_fieldnames(rows: Sequence[dict[str, Any]]) -> list[str]:
    fields: list[str] = []
    for row in rows:
        for key in row:
            if key not in fields:
                fields.append(key)
    return fields


def flatten_feature_row(sirna_id: str, feature: dict[str, Any]) -> dict[str, Any]:
    row = {"sirna_id": sirna_id}
    row.update(feature)
    return row


def write_workbook(
    workbook_path: Path,
    *,
    metadata: dict[str, Any],
    final_rows: Sequence[dict[str, Any]],
    feature_rows: Sequence[dict[str, Any]],
    garden_rows: Sequence[dict[str, Any]],
    mfe_rows: Sequence[dict[str, Any]],
    mismatch_rows: Sequence[dict[str, Any]],
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="1F6F78")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(bold=True, size=13)

    def add_rows(sheet_name: str, rows: Sequence[dict[str, Any]], preferred: Sequence[str] = ()) -> None:
        ws = wb.create_sheet(sheet_name)
        fields = [field for field in preferred if any(field in row for row in rows)]
        for field in collect_fieldnames(rows):
            if field not in fields:
                fields.append(field)
        if not fields:
            fields = ["note"]
            rows = [{"note": "No rows were produced."}]

        ws.append(fields)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for row in rows:
            ws.append([excel_value(row.get(field, "")) for field in fields])

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for index, field in enumerate(fields, start=1):
            max_width = len(str(field))
            for cell in ws.iter_cols(min_col=index, max_col=index, min_row=2, max_row=min(ws.max_row, 200)):
                for value_cell in cell:
                    max_width = max(max_width, min(len(str(value_cell.value or "")), 120))
            ws.column_dimensions[get_column_letter(index)].width = min(max(max_width + 2, 10), 42)

    readme = wb.create_sheet("README")
    readme["A1"] = "dsRNA to siRNA mismatch pipeline"
    readme["A1"].font = title_font
    for row_index, (key, value) in enumerate(metadata.items(), start=3):
        readme.cell(row_index, 1, key)
        readme.cell(row_index, 2, excel_value(value))
    readme.column_dimensions["A"].width = 32
    readme.column_dimensions["B"].width = 96

    add_rows(
        "Final_pairs",
        final_rows,
        preferred=[
            "dsrna_name",
            "sirna_id",
            "Transcript",
            "TStart",
            "TStop",
            "TSlice",
            "MFEratio",
            "AllenScore",
            "total_score",
            "TotalToleranceScore",
            "MismatchToleranceScore",
            "BulgeToleranceScore",
        ],
    )
    add_rows("siRNA_features", feature_rows, preferred=["sirna_id", "antisense_5_3", "total_score"])
    add_rows("garden_siRNAs", garden_rows, preferred=["sirna_id", "position_start", "position_end", "is_complete"])
    add_rows("MFE_pairs", mfe_rows, preferred=["Query", "Transcript", "MFEratio", "AllenScore"])
    add_rows("Mismatch_pairs", mismatch_rows, preferred=["Query", "Transcript", "TotalToleranceScore"])

    wb.save(workbook_path)


def excel_value(value: Any) -> Any:
    if isinstance(value, (str, int, float, bool)) or value is None:
        text = value
    else:
        text = json.dumps(value, sort_keys=True)
    if isinstance(text, str) and len(text) > 32000:
        return text[:31950] + "...[truncated]"
    return text


def score_features_for_duplexes(
    duplexes: Sequence[Any],
    temperature: float,
    accessibility: float,
    orf_status: str,
) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    rows: list[dict[str, Any]] = []
    by_id: dict[str, dict[str, Any]] = {}
    for duplex in duplexes:
        feature = predict_sirna_features(
            duplex.antisense_5_to_3,
            accessibility=accessibility,
            orf_status=orf_status,
            sense_overhang=duplex.sense_3prime_overhang,
            name=duplex.sirna_id,
            temperature=temperature,
            calculate_total_score=True,
        )
        feature["sirna_position_start"] = duplex.position_start
        feature["sirna_position_end"] = duplex.position_end
        by_id[duplex.sirna_id] = feature
        rows.append(flatten_feature_row(duplex.sirna_id, feature))
    return rows, by_id


def build_final_rows(
    dsrna: DsRNARecord,
    scored_mismatch_rows: Sequence[dict[str, Any]],
    feature_by_id: dict[str, dict[str, Any]],
    duplex_by_id: dict[str, Any],
) -> list[dict[str, Any]]:
    final_rows: list[dict[str, Any]] = []
    for row in scored_mismatch_rows:
        sirna_id = str(row.get("Query", ""))
        feature = feature_by_id.get(sirna_id, {})
        duplex = duplex_by_id.get(sirna_id)
        final: dict[str, Any] = {
            "dsrna_name": dsrna.name,
            "dsrna_length": len(dsrna.sequence_5_to_3),
            "sirna_id": sirna_id,
        }
        if duplex is not None:
            final.update(
                {
                    "sirna_number": duplex.number,
                    "dsrna_position_start": duplex.position_start,
                    "dsrna_position_end": duplex.position_end,
                    "antisense_5_to_3": duplex.antisense_5_to_3,
                    "sense_5_to_3": duplex.sense_5_to_3,
                    "sense_3_to_5": duplex.sense_3_to_5,
                }
            )
        for key, value in feature.items():
            final[key if key not in final else f"feature_{key}"] = value
        for key, value in row.items():
            final[key if key not in final else f"mfe_{key}"] = value
        final_rows.append(final)
    return final_rows


def write_detailed_mismatch_report(
    scored_rows: Sequence[dict[str, Any]],
    path: Path,
    *,
    position_mode: str,
    include_bulges: bool,
    bulge_stars: int,
) -> None:
    groups = []
    for index, row in enumerate(scored_rows, start=1):
        try:
            groups.append(
                pair_annotations_from_mfe_row(
                    row,
                    source_row=index,
                    position_mode=position_mode,
                    include_bulges=include_bulges,
                    bulge_stars=bulge_stars,
                )
            )
        except Exception:
            continue
    write_detailed_report(groups, path, bulge_stars=bulge_stars)


def process_dsrna_record(
    dsrna: DsRNARecord,
    transcript_fasta: Path,
    output_root: Path,
    args: argparse.Namespace,
) -> dict[str, Any]:
    dsrna_dir = output_root / safe_name(dsrna.name)
    if args.force and dsrna_dir.exists():
        shutil.rmtree(dsrna_dir)
    dsrna_dir.mkdir(parents=True, exist_ok=True)

    duplexes = list(
        generate_sirna_duplexes(
            dsrna,
            sirna_length=args.sirna_length,
            overhang_length=args.overhang_length,
        )
    )
    if args.limit_sirnas:
        duplexes = duplexes[: args.limit_sirnas]
    valid_duplexes = [duplex for duplex in duplexes if duplex.is_complete]
    duplex_by_id = {duplex.sirna_id: duplex for duplex in valid_duplexes}

    garden_rows = [duplex.to_dict() for duplex in duplexes]
    write_tsv(garden_rows, dsrna_dir / "01_garden_all_siRNAs.tsv")
    write_json(
        {
            "dsrna": asdict(dsrna),
            "sirna_length": args.sirna_length,
            "overhang_length": args.overhang_length,
            "duplex_count": len(duplexes),
            "valid_complete_duplex_count": len(valid_duplexes),
            "sirna_duplexes": garden_rows,
        },
        dsrna_dir / "01_garden_all_siRNAs.json",
    )

    valid_sirna_fasta = dsrna_dir / "02_valid_siRNAs.fasta"
    write_fasta_records(
        ((duplex.sirna_id, duplex.antisense_5_to_3) for duplex in valid_duplexes),
        valid_sirna_fasta,
    )

    feature_rows, feature_by_id = score_features_for_duplexes(
        valid_duplexes,
        temperature=args.temperature,
        accessibility=args.accessibility,
        orf_status=args.orf_status,
    )
    write_tsv(feature_rows, dsrna_dir / "03_siRNA_feature_scores.tsv")
    write_json(feature_rows, dsrna_dir / "03_siRNA_feature_scores.json")

    transcript_input = transcript_fasta
    limited_transcript_fasta = None
    if args.limit_transcripts:
        transcript_records = read_fasta_records(transcript_fasta, limit=args.limit_transcripts)
        limited_transcript_fasta = dsrna_dir / "00_limited_transcripts.fasta"
        write_fasta_records(transcript_records, limited_transcript_fasta)
        transcript_input = limited_transcript_fasta

    mfe_prefix = dsrna_dir / "04_MFE_ratio_pairs"
    run_mfe_analysis(
        srna_fasta=valid_sirna_fasta,
        transcript_fasta=transcript_input,
        output_prefix=mfe_prefix,
        mfe_ratio_cutoff=args.mfe_ratio_cutoff,
        sort_by=args.sort_by,
        top_n=args.top_n,
        rnaplex_path=args.rnaplex,
        write_png=args.write_png and not args.no_png,
        write_visualizations=not args.no_visualizations,
        visualization_limit=args.visualization_limit,
    )
    mfe_rows = read_tsv(mfe_prefix.with_suffix(".tsv"))

    mismatch_tsv = dsrna_dir / "05_mismatch_scored_pairs.tsv"
    pair_annotations_tsv = dsrna_dir / "05_mismatch_pair_annotations.tsv"
    alignment_annotations_tsv = dsrna_dir / "05_mismatch_alignment_annotations.tsv"
    vertical_annotations_tsv = dsrna_dir / "05_mismatch_vertical_annotations.tsv"
    score_file(
        input_path=mfe_prefix.with_suffix(".tsv"),
        output_path=mismatch_tsv,
        pair_table_output_path=pair_annotations_tsv,
        alignment_table_output_path=alignment_annotations_tsv,
        vertical_table_output_path=vertical_annotations_tsv,
        position_mode=args.position_mode,
        include_bulges=not args.ignore_bulges,
        bulge_stars=args.bulge_stars,
    )
    scored_mismatch_rows = read_tsv(mismatch_tsv)
    write_detailed_mismatch_report(
        scored_mismatch_rows,
        dsrna_dir / "05_mismatch_detailed_report.txt",
        position_mode=args.position_mode,
        include_bulges=not args.ignore_bulges,
        bulge_stars=args.bulge_stars,
    )

    final_rows = build_final_rows(dsrna, scored_mismatch_rows, feature_by_id, duplex_by_id)
    write_tsv(final_rows, dsrna_dir / "06_final_valid_siRNA_transcript_pairs.tsv")
    write_json(final_rows, dsrna_dir / "06_final_valid_siRNA_transcript_pairs.json")

    workbook_path = dsrna_dir / "06_final_valid_siRNA_transcript_pairs.xlsx"
    metadata = {
        "dsRNA": dsrna.name,
        "dsRNA length": len(dsrna.sequence_5_to_3),
        "transcript FASTA": str(transcript_fasta),
        "transcript FASTA used": str(transcript_input),
        "siRNA length": args.sirna_length,
        "overhang length": args.overhang_length,
        "complete siRNAs scored": len(valid_duplexes),
        "MFE ratio cutoff": args.mfe_ratio_cutoff,
        "valid siRNA-transcript pairs": len(final_rows),
        "boundary rule": "Incomplete duplexes containing X placeholders are retained in garden outputs but excluded downstream.",
    }
    write_workbook(
        workbook_path,
        metadata=metadata,
        final_rows=final_rows,
        feature_rows=feature_rows,
        garden_rows=garden_rows,
        mfe_rows=mfe_rows,
        mismatch_rows=scored_mismatch_rows,
    )

    summary = {
        "dsrna": dsrna.name,
        "output_dir": str(dsrna_dir),
        "workbook": str(workbook_path),
        "duplex_count": len(duplexes),
        "valid_complete_duplex_count": len(valid_duplexes),
        "mfe_pair_count": len(mfe_rows),
        "final_pair_count": len(final_rows),
        "limited_transcript_fasta": str(limited_transcript_fasta) if limited_transcript_fasta else "",
    }
    write_json(summary, dsrna_dir / "00_pipeline_summary.json")
    return summary


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate siRNAs from dsRNA FASTA and score siRNA-transcript mismatch pairs."
    )
    parser.add_argument("--dsrna-fasta", required=True, type=Path, help="Input dsRNA FASTA.")
    parser.add_argument("--transcripts-fasta", required=True, type=Path, help="Transcript FASTA.")
    parser.add_argument("--output-dir", type=Path, default=Path("outputs/dsRNA_to_siRNA_mismatch"))
    parser.add_argument("--sirna-length", type=int, default=DEFAULT_SIRNA_LENGTH)
    parser.add_argument("--overhang-length", type=int, default=DEFAULT_OVERHANG_LENGTH)
    parser.add_argument("--mfe-ratio-cutoff", type=float, default=0.70)
    parser.add_argument("--sort-by", choices=["mfe_ratio", "allen"], default="mfe_ratio")
    parser.add_argument("--top-n", type=int, default=0, help="Keep only top N MFE hits per dsRNA after sorting.")
    parser.add_argument("--rnaplex", default="RNAplex", help="RNAplex executable path.")
    parser.add_argument("--temperature", type=float, default=25)
    parser.add_argument("--accessibility", type=float, default=1.0)
    parser.add_argument(
        "--orf-status",
        choices=["ORF", "5_UTR", "3_UTR", "partial_ORF"],
        default="ORF",
    )
    parser.add_argument(
        "--position-mode",
        choices=["huang", "general", "none"],
        default="huang",
        help="Mismatch-sensitive position rule.",
    )
    parser.add_argument("--ignore-bulges", action="store_true")
    parser.add_argument("--bulge-stars", type=int, default=3)
    parser.add_argument("--limit-sirnas", type=int, default=0, help="Development/smoke-test limit.")
    parser.add_argument("--limit-transcripts", type=int, default=0, help="Development/smoke-test limit.")
    parser.add_argument("--no-visualizations", action="store_true", help="Skip per-hit MFE SVG visualizations.")
    parser.add_argument(
        "--visualization-limit",
        type=int,
        default=100,
        help="Maximum number of top sorted MFE hits to visualize. Default: 100.",
    )
    parser.add_argument("--write-visualizations", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--write-png", action="store_true", help="Write MFE summary PNG plot.")
    parser.add_argument("--no-png", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--force", action="store_true", help="Replace existing per-dsRNA output folders.")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    parser = build_arg_parser()
    args = parser.parse_args(argv)

    if not args.dsrna_fasta.exists():
        parser.error(f"dsRNA FASTA does not exist: {args.dsrna_fasta}")
    if not args.transcripts_fasta.exists():
        parser.error(f"Transcript FASTA does not exist: {args.transcripts_fasta}")

    try:
        dsrna_records = read_dsrna_fasta(args.dsrna_fasta)
    except DsRNAInputError as exc:
        parser.exit(2, f"Error: {exc}\n")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    summaries = [
        process_dsrna_record(record, args.transcripts_fasta, args.output_dir, args)
        for record in dsrna_records
    ]
    write_json(
        {
            "dsrna_fasta": str(args.dsrna_fasta),
            "transcripts_fasta": str(args.transcripts_fasta),
            "output_dir": str(args.output_dir),
            "dsrna_count": len(summaries),
            "summaries": summaries,
        },
        args.output_dir / "pipeline_summary.json",
    )
    for summary in summaries:
        print(
            f"{summary['dsrna']}: {summary['final_pair_count']} final pairs; "
            f"workbook={summary['workbook']}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
